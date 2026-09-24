"""Import Timetable Data from Excel (.xlsx) file into SQLite/SQLAlchemy Database.

Supports sheets:
- Departments: name, code, description
- Faculty: name, department, email, is_full_time, max_hours_per_week
- Courses: code, name, department, semester, credits, is_lab, default_periods_per_week, min_periods, faculty_email
- Sections: course_code, section_number, capacity, current_enrollment, periods_per_week, requires_lab
- Rooms: room_number, building, capacity, room_type, has_projector, has_computer
- TimeSlots: day_of_week, start_time, end_time, is_break, label
- Constraints: name, description, constraint_type, is_required, priority
"""
import sys
import os
import re
import argparse
from datetime import time, datetime
from typing import Dict, Any, List, Union, BinaryIO, Set, Optional
import io
import openpyxl
from sqlalchemy.orm import Session

# Add backend directory to sys.path so app imports work when run as script
BACKEND_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if BACKEND_DIR not in sys.path:
    sys.path.insert(0, BACKEND_DIR)

from app.database import SessionLocal, init_db, engine
from app.models.base import (
    Base,
    Department,
    Faculty,
    Course,
    Section,
    Room,
    TimeSlot,
    Constraint,
    FacultyAvailability,
    RoomAvailability,
    FacultyPreference,
    TimetableEntry,
)


def parse_time_value(val: Any) -> time:
    """Parse time from string (HH:MM or HH:MM:SS) or datetime.time object."""
    if isinstance(val, time):
        return val
    if isinstance(val, datetime):
        return val.time()
    if isinstance(val, str):
        val = val.strip()
        parts = val.split(":")
        if len(parts) >= 2:
            return time(int(parts[0]), int(parts[1]), int(parts[2]) if len(parts) > 2 else 0)
    raise ValueError(f"Unable to parse time from value: {val}")


def parse_bool(val: Any, default: bool = False) -> bool:
    """Parse boolean from various cell formats."""
    if val is None:
        return default
    if isinstance(val, bool):
        return val
    if isinstance(val, (int, float)):
        return bool(val)
    if isinstance(val, str):
        s = val.strip().lower()
        if s in ("true", "1", "yes", "y", "t"):
            return True
        if s in ("false", "0", "no", "n", "f"):
            return False
    return default


def parse_int(val: Any, default: int = 0) -> int:
    """Parse integer safely."""
    if val is None:
        return default
    try:
        return int(val)
    except (ValueError, TypeError):
        # Handle strings like "Semester 5" -> 5
        if isinstance(val, str):
            match = re.search(r'\d+', val)
            if match:
                return int(match.group())
        return default


def parse_semester(val: Any) -> Optional[int]:
    """Extract integer semester number (e.g. 5, '5', 'Semester 5', 'Sem 5', 'V')."""
    if val is None:
        return None
    if isinstance(val, int):
        return val
    if isinstance(val, str):
        val = val.strip()
        # Roman numerals map
        roman_map = {"i": 1, "ii": 2, "iii": 3, "iv": 4, "v": 5, "vi": 6, "vii": 7, "viii": 8}
        if val.lower() in roman_map:
            return roman_map[val.lower()]
        match = re.search(r'\d+', val)
        if match:
            return int(match.group())
    return None


def import_excel_data(
    file_source: Union[str, BinaryIO, bytes],
    db: Session,
    clear_existing: bool = False,
) -> Dict[str, Any]:
    """Import all timetable sheets from an Excel workbook into the database."""
    # Ensure tables exist
    init_db()

    # Load workbook
    if isinstance(file_source, bytes):
        wb = openpyxl.load_workbook(io.BytesIO(file_source), data_only=True)
    elif hasattr(file_source, "read"):
        content = file_source.read()
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    else:
        wb = openpyxl.load_workbook(file_source, data_only=True)

    counts = {
        "departments": 0,
        "faculty": 0,
        "courses": 0,
        "sections": 0,
        "rooms": 0,
        "time_slots": 0,
        "constraints": 0,
    }
    errors: List[str] = []
    warnings: List[str] = []

    try:
        if clear_existing:
            # Delete in order of foreign key constraints
            db.query(TimetableEntry).delete()
            db.query(FacultyAvailability).delete()
            db.query(RoomAvailability).delete()
            db.query(FacultyPreference).delete()
            db.query(Section).delete()
            db.query(Course).delete()
            db.query(Faculty).delete()
            db.query(Room).delete()
            db.query(TimeSlot).delete()
            db.query(Constraint).delete()
            db.query(Department).delete()
            db.commit()

        # Helper to get rows as dicts from a sheet
        def get_sheet_records(sheet_name: str) -> List[Dict[str, Any]]:
            # Case-insensitive sheet matching
            matched_name = next((s for s in wb.sheetnames if s.strip().lower() == sheet_name.strip().lower()), None)
            if not matched_name:
                return []
            ws = wb[matched_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows or len(rows) < 2:
                return []
            headers = [str(h).strip().lower() if h is not None else f"col_{idx}" for idx, h in enumerate(rows[0])]
            records = []
            for r in rows[1:]:
                if all(c is None for c in r):
                    continue
                record = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
                records.append(record)
            return records

        # 1. PRIMARY SHEET: "BE Subjects" / "Subjects"
        be_subject_records = get_sheet_records("BE Subjects")
        if not be_subject_records:
            be_subject_records = get_sheet_records("Subjects")
        if not be_subject_records:
            be_subject_records = get_sheet_records("BE_Subjects")

        # In-memory lookup maps for efficiency & uniqueness
        dept_name_map: Dict[str, Department] = {}
        for d in db.query(Department).all():
            if d.name:
                dept_name_map[d.name.strip().lower()] = d

        faculty_name_map: Dict[str, Faculty] = {}
        for f in db.query(Faculty).all():
            if f.name:
                faculty_name_map[f.name.strip().lower()] = f

        # Import from "BE Subjects" sheet if present
        if be_subject_records:
            for idx, row in enumerate(be_subject_records, start=2):
                dept_val = str(row.get("department") or row.get("dept") or "").strip()
                sem_val = str(row.get("semester") or row.get("sem") or "").strip()
                code_val = str(row.get("subject code") or row.get("subject_code") or row.get("code") or "").strip()
                name_val = str(row.get("subject name") or row.get("subject_name") or row.get("name") or "").strip()
                fac_name_val = str(row.get("faculty name") or row.get("faculty_name") or row.get("faculty") or "").strip()
                fac_initials_val = str(row.get("faculty initials") or row.get("faculty_initials") or row.get("initials") or "").strip() or None

                if not dept_val:
                    warnings.append(f"BE Subjects row {idx}: Missing Department. Skipping row.")
                    continue
                if not sem_val:
                    warnings.append(f"BE Subjects row {idx}: Missing Semester. Skipping row.")
                    continue

                # Ensure Department exists in database
                d_key = dept_val.lower()
                d_obj = dept_name_map.get(d_key)
                if not d_obj:
                    d_obj = Department(name=dept_val, code=None, description=None)
                    db.add(d_obj)
                    db.flush()
                    dept_name_map[d_key] = d_obj
                    counts["departments"] += 1

                # Handle Faculty record (create or update, deduplicate)
                fac_obj = None
                if fac_name_val:
                    fac_key = fac_name_val.lower()
                    fac_obj = faculty_name_map.get(fac_key)
                    if not fac_obj:
                        fac_obj = db.query(Faculty).filter(Faculty.name.ilike(fac_name_val)).first()

                    if not fac_obj:
                        fac_obj = Faculty(
                            name=fac_name_val,
                            initials=fac_initials_val,
                            department=d_obj.name,
                            department_id=d_obj.id,
                            is_full_time=True,
                            max_hours_per_week=20,
                        )
                        db.add(fac_obj)
                        db.flush()
                        faculty_name_map[fac_key] = fac_obj
                        counts["faculty"] += 1
                    else:
                        # Update missing initials or department if needed
                        if fac_initials_val and not fac_obj.initials:
                            fac_obj.initials = fac_initials_val
                        if d_obj.id and not fac_obj.department_id:
                            fac_obj.department_id = d_obj.id
                            fac_obj.department = d_obj.name
                        db.flush()
                        faculty_name_map[fac_key] = fac_obj

                if not code_val or not name_val:
                    warnings.append(f"BE Subjects row {idx}: Missing Subject Code or Name for Department '{dept_val}', Semester '{sem_val}'.")
                    continue

                # Uniqueness per (department_id, semester, code)
                c_obj = db.query(Course).filter(
                    Course.department_id == d_obj.id,
                    Course.semester == sem_val,
                    Course.code == code_val
                ).first()

                if not c_obj:
                    c_obj = Course(
                        code=code_val,
                        name=name_val,
                        department_id=d_obj.id,
                        semester=sem_val,
                        faculty_id=fac_obj.id if fac_obj else None,
                        credits=3,
                        default_periods_per_week=3,
                        min_periods=1,
                        is_lab=False,
                    )
                    db.add(c_obj)
                    db.flush()
                    counts["courses"] += 1
                else:
                    c_obj.name = name_val
                    if fac_obj:
                        c_obj.faculty_id = fac_obj.id
                    db.flush()
                    counts["courses"] += 1

                # Ensure Section exists for course
                sec_obj = db.query(Section).filter(Section.course_id == c_obj.id).first()
                if not sec_obj:
                    sec_obj = Section(
                        course_id=c_obj.id,
                        section_number="A",
                        capacity=60,
                        current_enrollment=30,
                        periods_per_week=c_obj.default_periods_per_week,
                        requires_lab=c_obj.is_lab,
                    )
                    db.add(sec_obj)
                    db.flush()
                    counts["sections"] += 1

        # 2. DEPARTMENTS SHEET (if dedicated sheet exists)
        dept_records = get_sheet_records("Departments")
        for idx, row in enumerate(dept_records, start=2):
            d_name = str(row.get("name") or "").strip()
            if not d_name:
                continue
            d_code = str(row.get("code") or "").strip() or None
            d_desc = str(row.get("description") or "").strip() or None

            d_obj = dept_name_map.get(d_name.lower())
            if not d_obj:
                d_obj = Department(name=d_name, code=d_code, description=d_desc)
                db.add(d_obj)
                db.flush()
                counts["departments"] += 1
                dept_name_map[d_name.lower()] = d_obj
            else:
                if d_code:
                    d_obj.code = d_code
                if d_desc:
                    d_obj.description = d_desc
                db.flush()

        # Helper function to find or auto-create a department
        def get_or_create_dept(dept_identifier: Optional[str]) -> Optional[Department]:
            if not dept_identifier:
                return None
            val = str(dept_identifier).strip()
            if not val:
                return None
            key = val.lower()
            if key in dept_name_map:
                return dept_name_map[key]
            # Auto-create if new
            new_dept = Department(name=val, code=None)
            db.add(new_dept)
            db.flush()
            counts["departments"] += 1
            dept_name_map[key] = new_dept
            return new_dept

        # 3. FACULTY SHEET (if present)
        faculty_email_map: Dict[str, Faculty] = {}
        faculty_name_map: Dict[str, Faculty] = {}
        for f in db.query(Faculty).all():
            if f.email:
                faculty_email_map[f.email.strip().lower()] = f
            if f.name:
                faculty_name_map[f.name.strip().lower()] = f

        faculty_records = get_sheet_records("Faculty")
        for idx, row in enumerate(faculty_records, start=2):
            name = str(row.get("name") or "").strip()
            if not name:
                warnings.append(f"Faculty sheet row {idx}: Missing faculty name.")
                continue
            email = str(row.get("email") or "").strip().lower() or f"fac_{len(faculty_name_map)+1}@school.edu"
            raw_dept = str(row.get("department") or "").strip() or None
            dept_obj = get_or_create_dept(raw_dept)
            is_full_time = parse_bool(row.get("is_full_time"), default=True)
            max_hours = parse_int(row.get("max_hours_per_week"), default=20)

            f_obj = faculty_email_map.get(email) or faculty_name_map.get(name.lower())
            if not f_obj:
                f_obj = Faculty(
                    name=name,
                    department=dept_obj.name if dept_obj else raw_dept,
                    department_id=dept_obj.id if dept_obj else None,
                    email=email,
                    is_full_time=is_full_time,
                    max_hours_per_week=max_hours,
                )
                db.add(f_obj)
                db.flush()
                counts["faculty"] += 1
            else:
                f_obj.name = name
                f_obj.department = dept_obj.name if dept_obj else raw_dept
                f_obj.department_id = dept_obj.id if dept_obj else f_obj.department_id
                f_obj.is_full_time = is_full_time
                f_obj.max_hours_per_week = max_hours
                db.flush()
                counts["faculty"] += 1

            faculty_email_map[email] = f_obj
            faculty_name_map[name.lower()] = f_obj

        # 4. COURSES SHEET (if dedicated Courses sheet exists and BE Subjects was not used)
        if not be_subject_records:
            course_records = get_sheet_records("Courses")
            for idx, row in enumerate(course_records, start=2):
                code = str(row.get("code") or "").strip().upper()
                name = str(row.get("name") or "").strip()
                if not code or not name:
                    warnings.append(f"Courses sheet row {idx}: Missing course code or name.")
                    continue

                credits = parse_int(row.get("credits"), default=3)
                is_lab = parse_bool(row.get("is_lab"), default=False)
                default_periods = parse_int(row.get("default_periods_per_week"), default=2 if is_lab else 3)
                min_periods = parse_int(row.get("min_periods"), default=1)

                raw_c_dept = str(row.get("department") or row.get("department_name") or row.get("department_code") or "").strip() or None
                dept_obj = get_or_create_dept(raw_c_dept)
                semester_val = str(row.get("semester") or "").strip() or None

                faculty_ref = str(row.get("faculty_email") or row.get("faculty") or row.get("faculty_name") or "").strip().lower()
                fac_obj = faculty_email_map.get(faculty_ref) or faculty_name_map.get(faculty_ref)
                faculty_id = fac_obj.id if fac_obj else None

                c_obj = db.query(Course).filter(
                    Course.department_id == (dept_obj.id if dept_obj else None),
                    Course.semester == semester_val,
                    Course.code == code
                ).first()

                if not c_obj:
                    c_obj = Course(
                        code=code,
                        name=name,
                        credits=credits,
                        faculty_id=faculty_id,
                        department_id=dept_obj.id if dept_obj else None,
                        semester=semester_val,
                        is_lab=is_lab,
                        default_periods_per_week=default_periods,
                        min_periods=min_periods,
                    )
                    db.add(c_obj)
                    db.flush()
                    counts["courses"] += 1
                else:
                    c_obj.name = name
                    c_obj.credits = credits
                    c_obj.faculty_id = faculty_id
                    c_obj.department_id = dept_obj.id if dept_obj else c_obj.department_id
                    c_obj.semester = semester_val if semester_val is not None else c_obj.semester
                    c_obj.is_lab = is_lab
                    c_obj.default_periods_per_week = default_periods
                    c_obj.min_periods = min_periods
                    db.flush()
                    counts["courses"] += 1

        # 5. SECTIONS SHEET (if present)
        section_records = get_sheet_records("Sections")
        for idx, row in enumerate(section_records, start=2):
            c_code = str(row.get("course_code") or row.get("course") or "").strip().upper()
            sec_num = str(row.get("section_number") or row.get("section") or "").strip()
            if not c_code or not sec_num:
                warnings.append(f"Sections sheet row {idx}: Missing course_code or section_number.")
                continue

            c_obj = db.query(Course).filter(Course.code == c_code).first()
            if not c_obj:
                warnings.append(f"Sections sheet row {idx}: Course '{c_code}' not found.")
                continue

            capacity = parse_int(row.get("capacity"), default=30)
            curr_enrollment = parse_int(row.get("current_enrollment"), default=0)
            periods_per_week = parse_int(row.get("periods_per_week"), default=c_obj.default_periods_per_week)
            requires_lab = parse_bool(row.get("requires_lab"), default=c_obj.is_lab)

            sec_obj = db.query(Section).filter(
                Section.course_id == c_obj.id,
                Section.section_number == sec_num
            ).first()

            if not sec_obj:
                sec_obj = Section(
                    course_id=c_obj.id,
                    section_number=sec_num,
                    capacity=capacity,
                    current_enrollment=curr_enrollment,
                    periods_per_week=periods_per_week,
                    requires_lab=requires_lab,
                )
                db.add(sec_obj)
                db.flush()
                counts["sections"] += 1
            else:
                sec_obj.capacity = capacity
                sec_obj.current_enrollment = curr_enrollment
                sec_obj.periods_per_week = periods_per_week
                sec_obj.requires_lab = requires_lab
                db.flush()
                counts["sections"] += 1

        # 6. ROOMS SHEET (if present)
        room_number_map: Dict[str, Room] = {}
        for r in db.query(Room).all():
            room_number_map[r.room_number.strip().lower()] = r

        room_records = get_sheet_records("Rooms")
        for idx, row in enumerate(room_records, start=2):
            room_num = str(row.get("room_number") or row.get("room") or "").strip()
            if not room_num:
                warnings.append(f"Rooms sheet row {idx}: Missing room_number.")
                continue

            building = str(row.get("building") or "").strip() or None
            capacity = parse_int(row.get("capacity"), default=40)
            room_type = str(row.get("room_type") or "lecture").strip().lower()
            has_projector = parse_bool(row.get("has_projector"), default=True)
            has_computer = parse_bool(row.get("has_computer"), default=(room_type == "lab"))
            
            raw_r_dept = str(row.get("department") or "").strip() or None
            dept_obj = get_or_create_dept(raw_r_dept)

            r_obj = room_number_map.get(room_num.lower())
            if not r_obj:
                r_obj = Room(
                    room_number=room_num,
                    building=building,
                    department_id=dept_obj.id if dept_obj else None,
                    capacity=capacity,
                    room_type=room_type,
                    has_projector=has_projector,
                    has_computer=has_computer,
                )
                db.add(r_obj)
                db.flush()
                counts["rooms"] += 1
            else:
                r_obj.building = building
                r_obj.department_id = dept_obj.id if dept_obj else r_obj.department_id
                r_obj.capacity = capacity
                r_obj.room_type = room_type
                r_obj.has_projector = has_projector
                r_obj.has_computer = has_computer
                db.flush()
                counts["rooms"] += 1

            room_number_map[room_num.lower()] = r_obj

        # 7. TIME SLOTS SHEET (if present)
        slot_records = get_sheet_records("TimeSlots")
        for idx, row in enumerate(slot_records, start=2):
            try:
                day_of_week = parse_int(row.get("day_of_week"), default=0)
                start_time = parse_time_value(row.get("start_time"))
                end_time = parse_time_value(row.get("end_time"))
                is_break = parse_bool(row.get("is_break"), default=False)
                label = str(row.get("label") or "").strip() or None

                ts_obj = db.query(TimeSlot).filter(
                    TimeSlot.day_of_week == day_of_week,
                    TimeSlot.start_time == start_time,
                    TimeSlot.end_time == end_time,
                ).first()

                if not ts_obj:
                    ts_obj = TimeSlot(
                        day_of_week=day_of_week,
                        start_time=start_time,
                        end_time=end_time,
                        is_break=is_break,
                        label=label,
                    )
                    db.add(ts_obj)
                    db.flush()
                    counts["time_slots"] += 1
                else:
                    ts_obj.is_break = is_break
                    ts_obj.label = label
                    db.flush()
                    counts["time_slots"] += 1
            except Exception as e:
                warnings.append(f"TimeSlots sheet row {idx}: {str(e)}")

        # 7. CONSTRAINTS
        constraint_records = get_sheet_records("Constraints")
        for idx, row in enumerate(constraint_records, start=2):
            name = str(row.get("name") or "").strip()
            if not name:
                continue
            desc = str(row.get("description") or "").strip() or None
            c_type = str(row.get("constraint_type") or "hard").strip().lower()
            is_req = parse_bool(row.get("is_required"), default=(c_type == "hard"))
            priority = parse_int(row.get("priority"), default=5)

            c_obj = db.query(Constraint).filter(Constraint.name == name).first()
            if not c_obj:
                c_obj = Constraint(
                    name=name,
                    description=desc,
                    constraint_type=c_type,
                    is_required=is_req,
                    priority=priority,
                )
                db.add(c_obj)
                db.flush()
                counts["constraints"] += 1
            else:
                c_obj.description = desc
                c_obj.constraint_type = c_type
                c_obj.is_required = is_req
                c_obj.priority = priority
                db.flush()
                counts["constraints"] += 1

        # Ensure default rooms exist if none in database
        if db.query(Room).count() == 0:
            default_rooms = [
                Room(room_number="Room 101", building="Main Building", capacity=60, room_type="lecture", has_projector=True, has_computer=False),
                Room(room_number="Room 102", building="Main Building", capacity=60, room_type="lecture", has_projector=True, has_computer=False),
                Room(room_number="Room 103", building="Main Building", capacity=60, room_type="lecture", has_projector=True, has_computer=False),
                Room(room_number="Room 104", building="Main Building", capacity=60, room_type="lecture", has_projector=True, has_computer=False),
                Room(room_number="Room 105", building="Main Building", capacity=60, room_type="lecture", has_projector=True, has_computer=False),
                Room(room_number="Lab 201", building="Science Block", capacity=40, room_type="lab", has_projector=True, has_computer=True),
                Room(room_number="Lab 202", building="Science Block", capacity=40, room_type="lab", has_projector=True, has_computer=True),
                Room(room_number="Lab 203", building="Science Block", capacity=40, room_type="lab", has_projector=True, has_computer=True),
            ]
            db.add_all(default_rooms)
            db.flush()
            counts["rooms"] += len(default_rooms)

        # Ensure default time slots exist if none in database
        if db.query(TimeSlot).count() == 0:
            slots = []
            slot_idx = 1
            slot_definitions = [
                (8, 9, False, "Period 1"),
                (9, 10, False, "Period 2"),
                (10, 11, False, "Period 3"),
                (11, 12, False, "Period 4"),
                (12, 13, True, "Lunch Break"),
                (13, 14, False, "Period 5"),
                (14, 15, False, "Period 6"),
                (15, 16, False, "Period 7"),
                (16, 17, False, "Period 8"),
            ]
            for day in range(5):  # Mon-Fri (0-4)
                for start_h, end_h, is_brk, label_prefix in slot_definitions:
                    lbl = label_prefix if is_brk else f"Slot {slot_idx}"
                    if not is_brk:
                        slot_idx += 1
                    slots.append(TimeSlot(
                        day_of_week=day,
                        start_time=time(start_h, 0),
                        end_time=time(end_h, 0),
                        is_break=is_brk,
                        label=lbl,
                    ))
            db.add_all(slots)
            db.flush()
            counts["time_slots"] += len(slots)

        # Ensure default constraints exist if none in database
        if db.query(Constraint).count() == 0:
            constraints = [
                Constraint(name="No Faculty Overlap", description="A faculty member cannot teach two courses at the same time", constraint_type="hard", is_required=True, priority=10),
                Constraint(name="No Section Overlap", description="A section cannot have two classes at the same time", constraint_type="hard", is_required=True, priority=9),
                Constraint(name="No Room Overlap", description="A room cannot host two classes at the same time", constraint_type="hard", is_required=True, priority=8),
                Constraint(name="Room Capacity", description="Room capacity must meet section demand", constraint_type="hard", is_required=True, priority=7),
                Constraint(name="Faculty Availability", description="Faculty must be available in assigned slots", constraint_type="hard", is_required=True, priority=6),
                Constraint(name="Lab Requirements", description="Lab sessions must use a lab room", constraint_type="hard", is_required=True, priority=5),
                Constraint(name="Break Periods", description="No classes during break slots", constraint_type="hard", is_required=True, priority=4),
                Constraint(name="Workload Balance", description="Faculty workload should be balanced", constraint_type="soft", is_required=False, priority=3),
                Constraint(name="Faculty Gaps", description="Minimize idle gaps between classes", constraint_type="soft", is_required=False, priority=2),
                Constraint(name="Student Gaps", description="Minimize gaps in student schedules", constraint_type="soft", is_required=False, priority=1),
            ]
            db.add_all(constraints)
            db.flush()
            counts["constraints"] += len(constraints)

        db.commit()
        return {
            "success": True,
            "message": f"Successfully imported data from Excel: {counts['departments']} departments, {counts['courses']} courses, {counts['faculty']} faculty, {counts['sections']} sections, {counts['rooms']} rooms, {counts['time_slots']} time slots, {counts['constraints']} constraints.",
            "counts": counts,
            "errors": errors,
            "warnings": warnings,
        }
    except Exception as e:
        db.rollback()
        return {
            "success": False,
            "message": f"Excel import failed: {str(e)}",
            "counts": counts,
            "errors": errors + [str(e)],
            "warnings": warnings,
        }


def main():
    parser = argparse.ArgumentParser(description="Import timetable data from Excel (.xlsx) file.")
    default_path = os.path.join(BACKEND_DIR, "data", "timetable.xlsx")
    parser.add_argument(
        "--file", "-f",
        default=default_path,
        help=f"Path to .xlsx file (default: {default_path})"
    )
    parser.add_argument(
        "--clear", "-c",
        action="store_true",
        help="Clear existing data before importing"
    )
    args = parser.parse_args()

    if not os.path.exists(args.file):
        print(f"Error: Excel file '{args.file}' not found.")
        sys.exit(1)

    print(f"\n==================================================")
    print(f"Excel Import Utility")
    print(f"File: {args.file}")
    print(f"Clear Existing: {args.clear}")
    print(f"==================================================\n")

    init_db()
    db = SessionLocal()
    try:
        res = import_excel_data(args.file, db, clear_existing=args.clear)
        if res["success"]:
            print(f"\nExcel Import Summary\n")
            for k, v in res["counts"].items():
                print(f"{k.replace('_', ' ').title()}: {v}")
            print(f"\nImported successfully.")
            print(f"Errors: {len(res['errors'])}")
            print(f"Warnings: {len(res.get('warnings', []))}")
            if res.get("warnings"):
                for w in res["warnings"]:
                    print(f"  [Warning] {w}")
        else:
            print(f"\nExcel Import Failed: {res['message']}")
            for err in res["errors"]:
                print(f"  [Error] {err}")
            sys.exit(1)
    finally:
        db.close()


if __name__ == "__main__":
    main()
