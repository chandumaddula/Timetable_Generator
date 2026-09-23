"""Generate a clean, structured sample timetable.xlsx workbook with multiple sheets."""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_excel_workbook(file_path: str):
    wb = openpyxl.Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    def style_sheet(ws, headers, rows):
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border
            ws.row_dimensions[1].height = 24

        for r_idx, row in enumerate(rows, start=2):
            ws.append(row)
            ws.row_dimensions[r_idx].height = 20
            for col_idx in range(1, len(row) + 1):
                cell = ws.cell(row=r_idx, column=col_idx)
                cell.font = Font(name="Calibri", size=10)
                cell.border = thin_border
                if isinstance(cell.value, (int, float, bool)):
                    cell.alignment = center_align
                elif str(cell.value).startswith("0") or str(cell.value).startswith("1") or str(cell.value).startswith("2"):
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # 1. Departments sheet
    dept_headers = ["name", "code", "description"]
    dept_rows = [
        ["Computer Science", "CSE", "Department of Computer Science & Engineering"],
        ["Electrical Engineering", "EEE", "Department of Electrical & Electronics Engineering"],
        ["Data Science", "DS", "Department of Data Science & AI"],
        ["Mathematics", "MATH", "Department of Mathematics & Computing"],
        ["Robotics", "ROB", "Department of Robotics & Automation"],
    ]
    ws_dept = wb.create_sheet(title="Departments")
    style_sheet(ws_dept, dept_headers, dept_rows)

    # 2. Faculty sheet
    faculty_headers = ["name", "department", "email", "is_full_time", "max_hours_per_week"]
    faculty_rows = [
        ["Prof. Alan Turing", "Computer Science", "turing@school.edu", True, 20],
        ["Prof. Grace Hopper", "Computer Science", "hopper@school.edu", True, 18],
        ["Prof. Donald Knuth", "Computer Science", "knuth@school.edu", True, 20],
        ["Prof. Barbara Liskov", "Computer Science", "liskov@school.edu", True, 18],
        ["Prof. Jeff Dean", "Computer Science", "dean@school.edu", True, 18],
        ["Prof. Michael Jordan", "Electrical Engineering", "jordan@school.edu", True, 22],
        ["Prof. Yann LeCun", "Electrical Engineering", "lecun@school.edu", True, 20],
        ["Prof. Andrew Ng", "Data Science", "ng@school.edu", True, 20],
        ["Prof. Fei-Fei Li", "Data Science", "feifeili@school.edu", True, 22],
        ["Prof. Ilya Sutskever", "Data Science", "sutskever@school.edu", True, 20],
        ["Prof. Demis Hassabis", "Data Science", "hassabis@school.edu", True, 20],
        ["Prof. Leslie Valiant", "Mathematics", "valiant@school.edu", True, 20],
        ["Prof. Shafi Goldwasser", "Mathematics", "goldwasser@school.edu", True, 18],
        ["Prof. Sebastian Thrun", "Robotics", "thrun@school.edu", True, 18],
    ]
    ws_faculty = wb.create_sheet(title="Faculty")
    style_sheet(ws_faculty, faculty_headers, faculty_rows)

    # 3. Courses sheet
    course_headers = ["code", "name", "department", "semester", "credits", "is_lab", "default_periods_per_week", "min_periods", "faculty_email"]
    course_rows = [
        # Computer Science
        ["CS101", "Intro to Programming", "Computer Science", 1, 3, True, 2, 1, "turing@school.edu"],
        ["CS102", "Digital Logic Design", "Computer Science", 1, 3, False, 3, 1, "hopper@school.edu"],
        ["CS201", "Data Structures", "Computer Science", 3, 3, True, 2, 1, "knuth@school.edu"],
        ["CS202", "Object Oriented Programming", "Computer Science", 3, 3, True, 2, 1, "liskov@school.edu"],
        ["CS301", "Database Systems", "Computer Science", 5, 3, False, 3, 1, "liskov@school.edu"],
        ["CS302", "Operating Systems", "Computer Science", 5, 3, False, 3, 1, "dean@school.edu"],
        ["CS303", "Computer Networks", "Computer Science", 5, 3, False, 3, 1, "turing@school.edu"],
        ["CS304", "Machine Learning", "Computer Science", 5, 3, True, 2, 1, "hopper@school.edu"],
        ["CS401", "Cloud Computing", "Computer Science", 7, 3, False, 3, 1, "dean@school.edu"],
        
        # Electrical Engineering
        ["EE101", "Basic Electrical Engineering", "Electrical Engineering", 1, 3, True, 2, 1, "jordan@school.edu"],
        ["EE201", "Circuit Theory", "Electrical Engineering", 3, 3, True, 2, 1, "lecun@school.edu"],
        ["EE301", "Power Systems", "Electrical Engineering", 5, 3, False, 3, 1, "jordan@school.edu"],
        ["EE302", "Control Systems", "Electrical Engineering", 5, 3, True, 2, 1, "lecun@school.edu"],
        ["EE303", "Electrical Machines", "Electrical Engineering", 5, 3, True, 2, 1, "jordan@school.edu"],
        
        # Data Science
        ["DS201", "Python for Data Science", "Data Science", 3, 3, True, 2, 1, "ng@school.edu"],
        ["DS301", "Big Data Analytics", "Data Science", 5, 3, True, 2, 1, "feifeili@school.edu"],
        ["DS302", "Applied Statistics", "Data Science", 5, 3, False, 3, 1, "sutskever@school.edu"],
        ["DS303", "Data Mining", "Data Science", 5, 3, True, 2, 1, "hassabis@school.edu"],
        
        # Mathematics
        ["MATH101", "Linear Algebra & Calculus", "Mathematics", 1, 4, False, 4, 1, "valiant@school.edu"],
        ["MATH201", "Discrete Mathematics", "Mathematics", 3, 4, False, 4, 1, "goldwasser@school.edu"],
        
        # Robotics
        ["ROB301", "Sensors & Actuators", "Robotics", 5, 3, True, 2, 1, "thrun@school.edu"],
        ["ROB302", "Robotic Kinematics", "Robotics", 5, 3, False, 3, 1, "thrun@school.edu"],
    ]
    ws_courses = wb.create_sheet(title="Courses")
    style_sheet(ws_courses, course_headers, course_rows)

    # 4. Sections sheet
    section_headers = ["course_code", "section_number", "capacity", "current_enrollment", "periods_per_week", "requires_lab"]
    section_rows = [
        ["CS101", "Sec-A1", 35, 30, 2, True],
        ["CS101", "Sec-A2", 35, 32, 2, True],
        ["CS102", "Sec-A1", 35, 30, 3, False],
        ["CS201", "Sec-B1", 30, 28, 2, True],
        ["CS201", "Sec-B2", 30, 25, 2, True],
        ["CS202", "Sec-B1", 35, 33, 2, True],
        ["CS301", "Sec-C1", 40, 35, 3, False],
        ["CS302", "Sec-C1", 40, 36, 3, False],
        ["CS303", "Sec-C1", 40, 38, 3, False],
        ["CS304", "Sec-C1", 40, 35, 2, True],
        ["CS401", "Sec-D1", 30, 28, 3, False],
        ["EE101", "Sec-A1", 35, 30, 2, True],
        ["EE201", "Sec-B1", 35, 32, 2, True],
        ["EE301", "Sec-C1", 35, 30, 3, False],
        ["EE302", "Sec-C1", 35, 30, 2, True],
        ["EE303", "Sec-C1", 35, 30, 2, True],
        ["DS201", "Sec-B1", 30, 28, 2, True],
        ["DS301", "Sec-C1", 30, 26, 2, True],
        ["DS302", "Sec-C1", 30, 28, 3, False],
        ["DS303", "Sec-C1", 30, 25, 2, True],
        ["MATH101", "Sec-A1", 45, 40, 4, False],
        ["MATH201", "Sec-B1", 40, 35, 4, False],
        ["ROB301", "Sec-C1", 30, 25, 2, True],
        ["ROB302", "Sec-C1", 30, 25, 3, False],
    ]
    ws_sections = wb.create_sheet(title="Sections")
    style_sheet(ws_sections, section_headers, section_rows)

    # 5. Rooms sheet
    room_headers = ["room_number", "building", "capacity", "room_type", "has_projector", "has_computer"]
    room_rows = [
        ["Room 101", "Engineering Building", 45, "lecture", True, False],
        ["Room 102", "Engineering Building", 45, "lecture", True, False],
        ["Lab 201", "Engineering Building", 35, "lab", True, True],
        ["Lab 202", "Engineering Building", 35, "lab", True, True],
        ["Room 301", "Science Building", 50, "lecture", True, False],
        ["Room 302", "Science Building", 50, "lecture", True, False],
        ["Lab 305", "Science Building", 30, "lab", True, True],
        ["Room 401", "Main Building", 60, "lecture", True, False],
        ["Room 402", "Main Building", 60, "lecture", True, False],
        ["Tutorial 1", "Main Building", 25, "tutorial", False, False],
    ]
    ws_rooms = wb.create_sheet(title="Rooms")
    style_sheet(ws_rooms, room_headers, room_rows)

    # 6. TimeSlots sheet
    slot_headers = ["day_of_week", "start_time", "end_time", "is_break", "label"]
    days = [0, 1, 2, 3, 4]  # Mon-Fri
    hours = [
        (8, 9, False, "Period 1"),
        (9, 10, False, "Period 2"),
        (10, 11, False, "Period 3"),
        (11, 12, False, "Period 4"),
        (12, 13, True, "Lunch Break"),
        (13, 14, False, "Period 5"),
        (14, 15, False, "Period 6"),
        (15, 16, False, "Period 7"),
        (16, 17, False, "Period 8"),
    ]
    slot_rows = []
    for day in days:
        for sh, eh, is_brk, lbl in hours:
            slot_rows.append([
                day,
                f"{sh:02d}:00",
                f"{eh:02d}:00",
                is_brk,
                lbl
            ])
    ws_slots = wb.create_sheet(title="TimeSlots")
    style_sheet(ws_slots, slot_headers, slot_rows)

    # 7. Constraints sheet
    constraint_headers = ["name", "description", "constraint_type", "is_required", "priority"]
    constraint_rows = [
        ["No Faculty Overlap", "A faculty member cannot teach two courses at the same time", "hard", True, 10],
        ["No Section Overlap", "A section cannot have two classes at the same time", "hard", True, 9],
        ["No Room Overlap", "A room cannot host two classes at the same time", "hard", True, 8],
        ["Room Capacity", "Room capacity must meet section demand", "hard", True, 7],
        ["Faculty Availability", "Faculty must be available in assigned slots", "hard", True, 6],
        ["Lab Requirements", "Lab sessions must use a lab room", "hard", True, 5],
        ["Break Periods", "No classes during break slots", "hard", True, 4],
        ["Workload Balance", "Faculty workload should be balanced", "soft", False, 3],
        ["Faculty Gaps", "Minimize idle gaps between classes", "soft", False, 2],
        ["Student Gaps", "Minimize gaps in student schedules", "soft", False, 1],
    ]
    ws_constraints = wb.create_sheet(title="Constraints")
    style_sheet(ws_constraints, constraint_headers, constraint_rows)

    os.makedirs(os.path.dirname(file_path), exist_ok=True)
    wb.save(file_path)
    print(f"Successfully generated timetable Excel workbook at: {file_path}")

if __name__ == "__main__":
    out_path = os.path.join(os.path.dirname(__file__), "..", "data", "timetable.xlsx")
    create_excel_workbook(out_path)
